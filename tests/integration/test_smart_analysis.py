"""Smart analysis integration using in-memory responses, never real credentials."""
from types import SimpleNamespace

import pytest

from src.core.batch_analysis import BatchAnalysisError, BatchAnalysisService
from src.core.provider_profile import ProviderProfile
from src.core.provider_protocol import ProviderCapabilities, ProviderResponse
from src.core.template_model import Template, TemplateField
from src.models.app_config import AppConfig
from src.services.retry_policy import RetryExecutionResult


TEXT = "明远科技的校园网络维护合同，签订于2026年5月。"
RESPONSE = {"understood": True, "language": "zh-CN", "parts": [
    {"role": "topic", "value": "校园网络维护合同", "evidence_source": "text", "evidence": "校园网络维护合同"},
    {"role": "entity", "value": "明远科技", "evidence_source": "text", "evidence": "明远科技"},
    {"role": "date", "value": "2026-05", "evidence_source": "text", "evidence": "2026年5月"},
]}


class Adapter:
    def __init__(self, response=RESPONSE, on_call=None):
        self.response = response
        self.on_call = on_call
        self.requests = []

    def invoke(self, request, *, cancellation_token=None):
        self.requests.append(request)
        if self.on_call:
            self.on_call()
        return RetryExecutionResult(True, ProviderResponse(fields=self.response), (), None, None)


def service(adapter, *, text=TEXT, template=None, config=None, vision=False, content_extractor=None, **kwargs):
    def no_template(*_):
        raise AssertionError("smart mode must not query templates")
    return BatchAnalysisService(
        profile_provider=lambda: ProviderProfile(id="fake", display_name="Fake", provider_id="openai", base_url="https://example.invalid/v1", model="test", capabilities=ProviderCapabilities(text=True, json_schema=True, vision=vision)),
        credential_provider=lambda _: "fake-not-a-secret",
        adapter_factory=lambda *_: adapter,
        config_provider=lambda: config or AppConfig(),
        template_provider=template or no_template,
        content_extractor=content_extractor or (lambda *_: (text, {"original_name": "not-evidence.txt"})),
        **kwargs,
    )


def source(tmp_path, suffix=".txt"):
    path = tmp_path / ("private-original" + suffix)
    path.write_bytes(b"fixture")
    return str(path)


def test_smart_bypasses_template_and_carries_stable_policy(tmp_path):
    adapter = Adapter()
    result = service(adapter)(source(tmp_path), "smart-1", naming_mode="smart")
    assert result == "2026-05_明远科技_校园网络维护合同.txt"
    request = adapter.requests[0]
    assert request.locale == "zh-CN"
    assert request.extracted_text == TEXT
    assert request.metadata["naming_policy_version"] == "smart-v1"
    assert request.metadata["naming_mode"] == "smart"
    assert "smart-v1" in request.requested_schema["description"]


@pytest.mark.parametrize("mode,version", [("unknown", "smart-v1"), ("smart", "future")])
def test_unknown_mode_and_policy_fail_without_provider(tmp_path, mode, version):
    adapter = Adapter()
    with pytest.raises(ValueError, match="^smart.error.unsupportedPolicy$"):
        service(adapter)(source(tmp_path), "smart", naming_mode=mode, naming_policy_version=version)
    assert adapter.requests == []


def test_empty_text_does_not_send_metadata_or_fallback_prompt_as_content(tmp_path):
    adapter = Adapter()
    with pytest.raises(ValueError, match="^smart.error.insufficientContent$"):
        service(adapter, text=" \n ")(source(tmp_path), "smart", naming_mode="smart")
    assert adapter.requests == []


def test_response_cancellation_checked_before_render(tmp_path):
    cancelled = [False]
    adapter = Adapter(on_call=lambda: cancelled.__setitem__(0, True))
    with pytest.raises(BatchAnalysisError, match="cancelled"):
        service(adapter)(source(tmp_path), "smart", naming_mode="smart", cancellation_token=lambda: cancelled[0])


def test_smart_explicit_language_reaches_policy(tmp_path):
    adapter = Adapter({"understood": True, "language": "en-US", "parts": [{"role": "topic", "value": "Campus network maintenance contract", "evidence_source": "text", "evidence": "校园网络维护合同"}]})
    result = service(adapter)(source(tmp_path), "smart", naming_mode="smart", output_language="en")
    assert result == "Campus network maintenance contract.txt"
    assert adapter.requests[0].locale == "en-US"


def test_smart_real_extracted_text_is_used_for_grounding(tmp_path):
    adapter = Adapter()
    with pytest.raises(ValueError, match="^smart.error.ungrounded$"):
        service(adapter, text="无关采购单")(source(tmp_path), "smart", naming_mode="smart")


def test_watch_and_default_call_remain_template_based(tmp_path):
    calls = []
    template = Template(name="fixture", pattern="{dcTitle}", fields=(TemplateField("dcTitle", "Title"),))
    def templates(template_id):
        calls.append(template_id)
        return template
    adapter = Adapter({"dcTitle": "原有模板命名"})
    pipeline = service(adapter, template=templates)
    path = source(tmp_path)
    assert pipeline(path, "template") == "原有模板命名.txt"
    result = pipeline.analyze(path, template_id="watch-template")
    assert result.suggested_name == "原有模板命名.txt"
    assert calls == ["", "watch-template"]
    assert "dcTitle" in adapter.requests[-1].requested_schema["properties"]


def test_smart_text_privacy_still_gates_provider(tmp_path):
    adapter = Adapter()
    with pytest.raises(BatchAnalysisError, match="Text cloud processing is disabled"):
        service(adapter, config=AppConfig(allow_text_cloud=False))(source(tmp_path), "smart", naming_mode="smart")
    assert adapter.requests == []


def test_smart_visual_requires_real_prepared_payload(tmp_path):
    adapter = Adapter()
    with pytest.raises(ValueError, match="^smart.error.insufficientContent$"):
        service(adapter, text="", vision=True, pdf_renderer=SimpleNamespace(render_file=lambda *a, **kw: SimpleNamespace(payloads=(), total_page_count=1, rendered_page_indices=())))(source(tmp_path, ".pdf"), "smart", naming_mode="smart")
    assert adapter.requests == []


def test_smart_invalid_output_language_uses_safe_error_id(tmp_path):
    with pytest.raises(ValueError, match="^smart.error.languageMismatch$"):
        service(Adapter())(source(tmp_path), "smart", naming_mode="smart", output_language="invalid-language")


def test_smart_prepared_visual_uses_local_extension_and_checks_vision_capability(tmp_path):
    from tests.integration.test_batch_analysis_service import _payload
    payload = _payload()
    image = SimpleNamespace(process_file=lambda *a, **kw: payload)
    response = {"understood": True, "language": "zh-CN", "parts": [{"role": "topic", "value": "红色自行车停靠路边", "evidence_source": "visual", "evidence": "观察到红色自行车"}]}
    adapter = Adapter(response)
    path = source(tmp_path, ".png")
    assert service(adapter, text="", vision=True, image_preprocessor=image)(path, "smart", naming_mode="smart") == "红色自行车停靠路边.png"
    assert adapter.requests[0].visual_payloads == (payload,)
    adapter.requests.clear()
    with pytest.raises(BatchAnalysisError, match="capability.vision_unavailable_image"):
        service(adapter, text="", image_preprocessor=image)(path, "smart", naming_mode="smart")
    assert adapter.requests == []


def media_source(tmp_path, suffix):
    from pathlib import Path
    from PIL import Image
    path = Path(source(tmp_path, suffix))
    if suffix == ".png":
        Image.new("RGB", (8, 8), "red").save(path)
    return str(path)


@pytest.mark.parametrize("suffix", [".png", ".mp4"])
def test_default_media_metadata_does_not_block_visual_subject(tmp_path, suffix):
    from src.core.batch_analysis import _default_extract
    from tests.integration.test_batch_analysis_service import _payload
    from pathlib import Path
    path = media_source(tmp_path, suffix)
    text, _ = _default_extract(Path(path), 2000)
    assert text.strip() and "private-original" in text
    payload = _payload()
    response = {"understood": True, "language": "zh-CN", "parts": [{
        "role": "topic", "value": "红色自行车停靠路边", "evidence_source": "visual",
        "evidence": "图像中有一辆红色自行车"}]}
    adapter = Adapter(response)
    pipeline = service(adapter, vision=True, content_extractor=_default_extract,
                       image_preprocessor=SimpleNamespace(process_file=lambda *a, **kw: payload),
                       video_extractor=SimpleNamespace(extract_file=lambda *a, **kw: SimpleNamespace(payloads=(payload,))))
    assert pipeline(path, "smart", naming_mode="smart") == "红色自行车停靠路边" + suffix
    assert "private-original" not in adapter.requests[0].extracted_text
    assert adapter.requests[0].visual_payloads == (payload,)


@pytest.mark.parametrize("suffix", [".png", ".mp4"])
def test_default_media_metadata_cannot_be_text_evidence(tmp_path, suffix):
    from src.core.batch_analysis import _default_extract
    from tests.integration.test_batch_analysis_service import _payload
    payload = _payload()
    response = {"understood": True, "language": "zh-CN", "parts": [{
        "role": "topic", "value": "校园网络维护合同", "evidence_source": "text",
        "evidence": "private-original"}]}
    pipeline = service(Adapter(response), vision=True, content_extractor=_default_extract,
                       image_preprocessor=SimpleNamespace(process_file=lambda *a, **kw: payload),
                       video_extractor=SimpleNamespace(extract_file=lambda *a, **kw: SimpleNamespace(payloads=(payload,))))
    with pytest.raises(ValueError, match="^smart.error.ungrounded$"):
        pipeline(media_source(tmp_path, suffix), "smart", naming_mode="smart")


def test_video_metadata_without_prepared_frames_fails_before_provider(tmp_path):
    from src.core.batch_analysis import _default_extract
    adapter = Adapter()
    pipeline = service(adapter, vision=True, content_extractor=_default_extract,
                       video_extractor=SimpleNamespace(extract_file=lambda *a, **kw: SimpleNamespace(payloads=())))
    with pytest.raises(ValueError, match="^smart.error.insufficientContent$"):
        pipeline(source(tmp_path, ".mp4"), "smart", naming_mode="smart")
    assert adapter.requests == []


def test_empty_workbook_sheet_labels_are_not_understanding(tmp_path):
    from openpyxl import Workbook
    from src.core.batch_analysis import _default_extract
    path = tmp_path / "blank.xlsx"
    book = Workbook()
    book.active.title = "Campus network contract"
    book.create_sheet("Sheet2")
    book.save(path)
    book.close()
    text, _ = _default_extract(path, 2000)
    assert "Campus network contract" in text
    adapter = Adapter({"understood": True, "language": "en-US", "parts": [{
        "role": "topic", "value": "Campus network contract", "evidence_source": "text",
        "evidence": "Campus network contract"}]})
    pipeline = service(adapter, content_extractor=_default_extract)
    with pytest.raises(ValueError, match="^smart.error.insufficientContent$"):
        pipeline(str(path), "blank", naming_mode="smart", output_language="en")
    assert adapter.requests == []


def test_workbook_with_actual_cells_remains_available(tmp_path):
    from openpyxl import Workbook
    from src.core.batch_analysis import _default_extract
    path = tmp_path / "data.xlsx"
    book = Workbook()
    book.active["A1"] = TEXT
    book.save(path)
    book.close()
    adapter = Adapter()
    result = service(adapter, content_extractor=_default_extract)(str(path), "cells", naming_mode="smart")
    assert result == "2026-05_明远科技_校园网络维护合同.xlsx"
    assert len(adapter.requests) == 1
