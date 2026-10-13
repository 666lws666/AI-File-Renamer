"""Bounded, cancellable extraction of synthetic local documents."""
import builtins
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from loguru import logger

from src.core import content_extractor as subject
from src.models.file_item import FileItem, FileStatus


def item(tmp_path, data=b"", suffix=".txt"):
    path = tmp_path / ("private-marker" + suffix)
    path.write_bytes(data)
    return FileItem(original_path=path)


@pytest.mark.parametrize("encoding", ["utf-8", "gbk", "utf-16"])
def test_bounded_text_preserves_cjk_and_terminal_multibyte(tmp_path, encoding):
    source = item(tmp_path, ("报告季度增长" * 10000).encode(encoding))
    subject.extract(source, 7)
    assert source.extracted_text == "报告季度增长报"
    assert source.metadata["text_truncated"] is True
    assert source.metadata["coverage_reason"] == "character_limit"
    assert "private-marker" not in repr(source.metadata)


def test_text_read_is_bounded_before_allocation(tmp_path, monkeypatch):
    source = item(tmp_path, b"x" * 1_000_000)
    calls = []
    real_open = builtins.open

    class ReadSpy:
        def __init__(self, stream): self.stream = stream
        def __enter__(self): return self
        def __exit__(self, *args): self.stream.close()
        def read(self, size=-1):
            calls.append(size)
            assert 0 < size <= 4096
            return self.stream.read(size)
        def __getattr__(self, name): return getattr(self.stream, name)

    monkeypatch.setattr(subject, "open", lambda *a, **k: ReadSpy(real_open(*a, **k)), raising=False)
    subject.extract(source, 31)
    assert source.status == FileStatus.EXTRACTED
    assert source.extracted_text == "x" * 31
    assert calls and sum(calls) <= 4 * (31 + 1) + 4


@pytest.mark.parametrize("text,limit,truncated", [("abc", 3, False), ("abcd", 3, True), ("", 0, False), ("a", 0, True)])
def test_exact_limit_is_not_assumed_truncated(tmp_path, text, limit, truncated):
    source = item(tmp_path, text.encode())
    subject.extract(source, limit)
    assert source.extracted_text == text[:limit]
    assert source.metadata["text_truncated"] is truncated
    assert source.metadata["coverage_reason"] == ("character_limit" if truncated else "complete")


@pytest.mark.parametrize("data,limit,total,reason", [
    (b"tiny text", 40, 1, "complete"), (b"", 0, 0, "complete"),
    (b"abc", 3, 1, "complete"), (b"abcd", 3, 1, "character_limit"),
    (b"x" * 20, 3, None, "character_limit"),
])
def test_text_eof_publishes_only_observed_section_total(tmp_path, data, limit, total, reason):
    from src.core.managed_watch_metadata import ManagedWatchMetadata, auto_apply_review_reason
    from src.services.local_ai.preparation_protocol import PreparationCoverage
    text, metadata = subject._extract_text(item(tmp_path, data).original_path, max_chars=limit)
    assert metadata.get("sections_total") == total
    assert metadata["coverage_reason"] == reason
    coverage = PreparationCoverage("text", text_source="extracted_text" if text else "none",
        sections_processed=metadata["sections_processed"], sections_total=metadata.get("sections_total"),
        text_truncated=metadata["text_truncated"], text_limit_reason=reason)
    warnings = () if reason == "complete" else ("text_character_limit",)
    record = ManagedWatchMetadata("text", "balanced", "balanced", coverage, warnings)
    assert auto_apply_review_reason(record) == (
        "partial_content" if reason != "complete" else "" if text else "coverage_unavailable")
    unknown = PreparationCoverage("text", sections_processed=1)
    assert auto_apply_review_reason(ManagedWatchMetadata("text", "balanced", "balanced", unknown)) == "coverage_unavailable"


@pytest.mark.parametrize("data,chars,byte_limit,expected_reason", [
    (b"tiny", 4, 100, ""), (b"", 4, 100, "coverage_unavailable"),
    (b"tiny-more", 4, 100, "partial_content"),
    ("雨水花園".encode("utf-8"), 4, 6, "partial_content"),
])
def test_actual_worker_inherits_text_section_total_without_bypassing_byte_limit(
        tmp_path, data, chars, byte_limit, expected_reason):
    from src.core.managed_watch_metadata import ManagedWatchMetadata, auto_apply_review_reason
    from src.services.local_ai.preparation_protocol import PreparationRequest, LocalPreparationLimits
    from src.services.local_ai.preparation_worker import prepare_request
    source = item(tmp_path, data)
    request = PreparationRequest(str(source.original_path),
        LocalPreparationLimits(max_text_chars=chars, max_text_utf8_bytes=byte_limit), "balanced")
    result = prepare_request(request)
    assert result.coverage.sections_total == (1 if data else 0)
    if byte_limit == 6:
        assert result.coverage.text_limit_reason == "utf8_limit"
        assert result.warning_codes == ("text_utf8_limit",)
    assert auto_apply_review_reason(ManagedWatchMetadata(
        "text", "balanced", "balanced", result.coverage, result.warning_codes)) == expected_reason


def test_cancelled_text_never_retries_an_encoding(tmp_path, monkeypatch):
    source = item(tmp_path, "报告".encode("gbk"))
    opens = []
    real_open = builtins.open
    monkeypatch.setattr(subject, "open", lambda *a, **k: (opens.append(a), real_open(*a, **k))[1], raising=False)
    subject.extract(source, cancellation_token=lambda: True)
    assert source.status == FileStatus.FAILED
    assert source.error_message == "extraction_cancelled"
    assert opens == []


def test_text_cancellation_after_read_never_publishes_complete_coverage(tmp_path, monkeypatch):
    source = item(tmp_path, b"tiny text")
    real_open = builtins.open
    did_read = False
    class ReadThenCancel:
        def __init__(self, stream): self.stream = stream
        def __enter__(self): return self
        def __exit__(self, *args): self.stream.close()
        def read(self, size):
            nonlocal did_read
            result = self.stream.read(size)
            did_read = True
            return result
    monkeypatch.setattr(subject, "open", lambda *a, **k: ReadThenCancel(real_open(*a, **k)), raising=False)
    subject.extract(source, cancellation_token=lambda: did_read, redacted_diagnostics=True)
    assert did_read and source.status == FileStatus.FAILED
    assert source.error_message == "extraction_cancelled"
    assert "sections_total" not in source.metadata


@pytest.mark.parametrize("suffix", [".txt", ".pdf", ".docx", ".pptx", ".xlsx", ".png", ".mp4"])
def test_redacted_failures_never_emit_filename_or_exception(tmp_path, monkeypatch, suffix):
    source = item(tmp_path, suffix=suffix)
    helper = {".txt": "text", ".pdf": "pdf", ".docx": "docx", ".pptx": "pptx", ".xlsx": "xlsx", ".png": "image", ".mp4": "video"}[suffix]
    def fail(*args, **kwargs): raise ValueError("private-marker secret-content C:/private/path")
    monkeypatch.setattr(subject, "_extract_" + helper, fail)
    messages = []
    sink = logger.add(lambda message: messages.append(str(message)))
    try:
        subject.extract(source, redacted_diagnostics=True)
    finally:
        logger.remove(sink)
    assert source.status == FileStatus.FAILED
    assert source.error_message == "extraction_failed"
    assert "private-marker" not in "".join(messages)
    assert "secret-content" not in "".join(messages)
    assert "C:/private/path" not in "".join(messages)


def test_pdf_exact_cap_and_page_coverage(tmp_path, monkeypatch):
    import PyPDF2
    source = item(tmp_path, suffix=".pdf")
    visits = []
    pages = [SimpleNamespace(extract_text=lambda i=i: (visits.append(i), "abc")[1]) for i in range(3)]
    monkeypatch.setattr(PyPDF2, "PdfReader", lambda _: SimpleNamespace(pages=pages, metadata={"/Author": "Author"}))
    subject.extract(source, 3, max_pdf_pages=1)
    assert visits == [0]
    assert source.extracted_text == "abc"
    assert source.metadata["pages_processed"] == 1
    assert source.metadata["pages"] == 3
    assert source.metadata["coverage_reason"] == "section_limit"
    assert source.metadata["author"] == "Author"


def test_real_mixed_pdf_distinguishes_visited_from_text_pages(tmp_path):
    from tests.integration.test_local_analysis_routing import synthetic_pdf
    source = item(tmp_path, suffix=".pdf")
    synthetic_pdf(source.original_path, text_pages=1, scan_pages=1)
    subject.extract(source, 4000, max_pdf_pages=2, redacted_diagnostics=True)
    assert source.status is FileStatus.EXTRACTED
    assert source.extracted_text.strip() == "Synthetic cover title"
    assert source.metadata["pages_processed"] == source.metadata["sections_processed"] == 2
    assert source.metadata["pages"] == source.metadata["sections_total"] == 2
    assert source.metadata["text_page_indices"] == (0,)
    assert source.metadata["coverage_reason"] == "complete"


@pytest.mark.parametrize("limit,indices,reason", [(0, (), "character_limit"), (5, (0,), "character_limit"), (21, (0,), "character_limit")])
def test_pdf_text_page_facts_only_count_nonwhitespace_returned_body(tmp_path, limit, indices, reason):
    from tests.integration.test_local_analysis_routing import synthetic_pdf
    source = item(tmp_path, suffix=".pdf")
    synthetic_pdf(source.original_path, text_pages=2, scan_pages=0)
    subject.extract(source, limit, max_pdf_pages=2, redacted_diagnostics=True)
    assert source.metadata["text_page_indices"] == indices
    assert source.metadata["coverage_reason"] == reason


@pytest.mark.parametrize("requested,expected", [(None, 100), (200, 200), (10000, 200)])
def test_real_pdf_cloud_allowance_has_absolute_200_cap_and_default_100(tmp_path, requested, expected):
    from tests.integration.test_local_analysis_routing import synthetic_pdf
    source = item(tmp_path, suffix=".pdf")
    synthetic_pdf(source.original_path, text_pages=201, scan_pages=0)
    subject.extract(source, 10000, max_pdf_pages=requested, redacted_diagnostics=True)
    assert source.metadata["pages_processed"] == expected
    assert source.metadata["text_page_indices"] == tuple(range(expected))
    assert source.metadata["coverage_reason"] == "section_limit"


def test_pdf_cancellation_checked_after_load_and_page(tmp_path, monkeypatch):
    import PyPDF2
    source = item(tmp_path, suffix=".pdf")
    visits = []
    streams = []
    def load(stream):
        streams.append(stream)
        return SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: (visits.append(1), "abc")[1])] * 10, metadata={})
    monkeypatch.setattr(PyPDF2, "PdfReader", load)
    subject.extract(source, cancellation_token=lambda: bool(visits))
    assert source.error_message == "extraction_cancelled"
    assert visits == [1]
    assert streams[0].closed


def test_office_preflight_rejects_expanded_member_before_parser(tmp_path, monkeypatch):
    import docx
    source = item(tmp_path, suffix=".docx")
    with ZipFile(source.original_path, "w") as package:
        package.writestr("word/document.xml", "x" * 100)
    monkeypatch.setattr(subject, "_MAX_OOXML_MEMBER_BYTES", 50, raising=False)
    called = []
    monkeypatch.setattr(docx, "Document", lambda _: called.append(True))
    subject.extract(source, redacted_diagnostics=True)
    assert called == []
    assert source.error_message == "document_limit"


def test_xlsx_cancellation_closes_workbook(tmp_path, monkeypatch):
    import openpyxl
    source = item(tmp_path, suffix=".xlsx")
    with ZipFile(source.original_path, "w") as package: package.writestr("fixture", "synthetic")
    visited = []
    closed = []
    class Workbook:
        sheetnames = ["one"]
        def __getitem__(self, name):
            def rows(**kwargs):
                for i in range(100):
                    visited.append(i)
                    yield ("cell",)
            return SimpleNamespace(iter_rows=rows, max_row=100, max_column=1)
        def close(self): closed.append(True)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: Workbook())
    subject.extract(source, cancellation_token=lambda: bool(visited))
    assert visited == [0]
    assert closed == [True]
    assert source.error_message == "extraction_cancelled"


@pytest.mark.parametrize("suffix", [".m4v", ".ts", ".mts", ".m2ts", ".3gp", ".3g2", ".mxf", ".mpg", ".mpeg", ".ogv"])
def test_supported_video_dispatch_remains_metadata_only(tmp_path, suffix):
    source = item(tmp_path, suffix=suffix)
    subject.extract(source)
    assert source.status == FileStatus.EXTRACTED
    assert source.metadata["content_source"] == "metadata_only"


@pytest.mark.parametrize("suffix", [".docx", ".pptx", ".xlsx"])
def test_real_office_character_budget_and_coverage(tmp_path, suffix):
    source = item(tmp_path, suffix=suffix)
    if suffix == ".docx":
        from docx import Document
        doc = Document()
        for _ in range(5): doc.add_paragraph("Synthetic report text")
        doc.save(source.original_path)
    elif suffix == ".pptx":
        from pptx import Presentation
        doc = Presentation()
        for _ in range(5): doc.slides.add_slide(doc.slide_layouts[0]).shapes.title.text = "Synthetic report text"
        doc.save(source.original_path)
    else:
        from openpyxl import Workbook
        doc = Workbook()
        for _ in range(5): doc.active.append(["Synthetic report text"])
        doc.save(source.original_path)
        doc.close()
    subject.extract(source, 8, redacted_diagnostics=True)
    assert source.status == FileStatus.EXTRACTED
    assert len(source.extracted_text) == 8
    assert source.metadata["text_truncated"] is True
    assert source.metadata["coverage_reason"] == "character_limit"
    assert source.metadata["sections_processed"] <= 1


def test_absurd_character_limit_is_rejected_without_read(tmp_path):
    source = item(tmp_path, b"synthetic")
    subject.extract(source, 10 ** 12, redacted_diagnostics=True)
    assert source.status == FileStatus.FAILED
    assert source.error_message == "invalid_extraction_limit"


def test_image_is_closed_on_success(tmp_path, monkeypatch):
    from PIL import Image
    source = item(tmp_path, suffix=".png")
    closed = []
    fake = SimpleNamespace(size=(10, 20), format="PNG", mode="RGB", _getexif=lambda: {}, close=lambda: closed.append(True))
    monkeypatch.setattr(Image, "open", lambda _: fake)
    subject.extract(source)
    assert source.status == FileStatus.EXTRACTED
    assert closed == [True]


@pytest.mark.parametrize("limit_name,limit", [("_MAX_OOXML_MEMBERS", 1), ("_MAX_OOXML_EXPANDED_BYTES", 15), ("_MAX_DOCUMENT_BYTES", 10)])
def test_office_preflight_all_fixed_limits_before_parser(tmp_path, monkeypatch, limit_name, limit):
    import docx
    source = item(tmp_path, suffix=".docx")
    with ZipFile(source.original_path, "w") as package:
        package.writestr("one", "x" * 10)
        package.writestr("two", "x" * 10)
    monkeypatch.setattr(subject, limit_name, limit)
    called = []
    monkeypatch.setattr(docx, "Document", lambda _: called.append(True))
    subject.extract(source, redacted_diagnostics=True)
    assert source.error_message == "document_limit"
    assert called == []


def test_xlsx_bounds_columns_and_reports_unexamined_columns(tmp_path, monkeypatch):
    import openpyxl
    source = item(tmp_path, suffix=".xlsx")
    with ZipFile(source.original_path, "w") as package: package.writestr("fixture", "synthetic")
    requested = []
    closed = []
    class Workbook:
        sheetnames = ["one"]
        def __getitem__(self, name):
            def rows(**kwargs):
                requested.append(kwargs)
                yield ("cell",)
            return SimpleNamespace(iter_rows=rows, max_row=1, max_column=10 ** 12)
        def close(self): closed.append(True)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: Workbook())
    subject.extract(source)
    assert source.status == FileStatus.EXTRACTED
    assert requested[0]["max_col"] <= 100
    assert requested[0]["max_row"] <= 200
    assert source.metadata["coverage_reason"] == "section_limit"
    assert "sections_total" not in source.metadata
    assert closed == [True]


@pytest.mark.parametrize("suffix", [".docx", ".pptx"])
def test_cancel_immediately_after_office_load_closes_source(tmp_path, monkeypatch, suffix):
    import docx
    import pptx
    source = item(tmp_path, suffix=suffix)
    with ZipFile(source.original_path, "w") as package: package.writestr("fixture", "synthetic")
    streams = []
    def loader(stream):
        streams.append(stream)
        return object()  # Accessing document properties must never happen.
    monkeypatch.setattr(docx if suffix == ".docx" else pptx, "Document" if suffix == ".docx" else "Presentation", loader)
    subject.extract(source, cancellation_token=lambda: bool(streams))
    assert source.error_message == "extraction_cancelled"
    assert streams[0].closed


def test_image_closes_when_cancelled_after_open(tmp_path, monkeypatch):
    from PIL import Image
    source = item(tmp_path, suffix=".png")
    opened, closed = [], []
    def load(_):
        opened.append(True)
        return SimpleNamespace(close=lambda: closed.append(True))
    monkeypatch.setattr(Image, "open", load)
    subject.extract(source, cancellation_token=lambda: bool(opened))
    assert source.error_message == "extraction_cancelled"
    assert closed == [True]


def test_plaintext_preserves_universal_newline_semantics(tmp_path):
    source = item(tmp_path, b"a\r\nb\rc\n")
    subject.extract(source, 6)
    assert source.extracted_text == "a\nb\nc\n"
    assert source.metadata["text_truncated"] is False


def test_xlsx_source_is_closed_even_when_parser_load_fails(tmp_path, monkeypatch):
    import openpyxl
    source = item(tmp_path, suffix=".xlsx")
    with ZipFile(source.original_path, "w") as package: package.writestr("fixture", "synthetic")
    streams = []
    def fail(stream, **kwargs):
        streams.append(stream)
        raise ValueError("synthetic-parser-failure")
    monkeypatch.setattr(openpyxl, "load_workbook", fail)
    subject.extract(source, redacted_diagnostics=True)
    assert source.error_message == "extraction_failed"
    assert not isinstance(streams[0], str)
    assert streams[0].closed


def test_xlsx_does_not_count_fabricated_empty_rows(tmp_path):
    from openpyxl import Workbook
    source = item(tmp_path, suffix=".xlsx")
    wb = Workbook()
    wb.active.append(["synthetic"])
    wb.save(source.original_path)
    wb.close()
    subject.extract(source)
    assert source.metadata["rows_processed"] == 1
    assert source.metadata["coverage_reason"] == "complete"


def test_image_preserves_existing_exif_description_metadata(tmp_path, monkeypatch):
    from PIL import Image
    source = item(tmp_path, suffix=".png")
    description = "synthetic " * 40
    fake = SimpleNamespace(size=(10, 20), format="PNG", mode="RGB", _getexif=lambda: {270: description}, close=lambda: None)
    monkeypatch.setattr(Image, "open", lambda _: fake)
    subject.extract(source)
    assert source.metadata["exif_ImageDescription"] == description
    assert source.extracted_text == "描述: " + description


@pytest.mark.parametrize("suffix", [".docx", ".pptx"])
def test_blank_document_sections_are_bounded_and_coverage_is_partial(tmp_path, monkeypatch, suffix):
    source = item(tmp_path, suffix=suffix)
    if suffix == ".docx":
        from docx import Document
        doc = Document()
        for _ in range(3): doc.add_paragraph("")
        doc.save(source.original_path)
        monkeypatch.setattr(subject, "_MAX_SECTIONS", 2)
    else:
        from pptx import Presentation
        doc = Presentation()
        for _ in range(3): doc.slides.add_slide(doc.slide_layouts[6])
        doc.save(source.original_path)
        monkeypatch.setattr(subject, "_MAX_SLIDES", 2)
    subject.extract(source)
    assert source.status == FileStatus.EXTRACTED
    assert source.extracted_text == ""
    assert source.metadata["sections_processed"] == 2
    assert source.metadata["sections_total"] == 3
    assert source.metadata["coverage_reason"] == "section_limit"


@pytest.mark.parametrize("suffix", [".docx", ".pptx"])
def test_exact_final_office_section_does_not_report_truncation(tmp_path, suffix):
    source = item(tmp_path, suffix=suffix)
    if suffix == ".docx":
        from docx import Document
        doc = Document()
        doc.add_paragraph("abc")
    else:
        from pptx import Presentation
        doc = Presentation()
        doc.slides.add_slide(doc.slide_layouts[0]).shapes.title.text = "abc"
    doc.save(source.original_path)
    subject.extract(source, 3)
    assert source.extracted_text == "abc"
    assert source.metadata["text_truncated"] is False
    assert source.metadata["coverage_reason"] == "complete"
